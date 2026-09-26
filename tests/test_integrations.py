import os
from pathlib import Path
import pytest
import pymupdf as fitz
from engine import add_text
import advanced


def test_bundled_turkish_ocr(tmp_path):
    with fitz.open() as doc:
        page = doc.new_page()
        add_text(page, fitz.Rect(40, 40, 540, 150), 'OCR test Istanbul 12345', 24)
        output = tmp_path / 'ocr.pdf'
        advanced.ocr_pdf(doc.tobytes(), output)
    with fitz.open(output) as result:
        assert '12345' in result[0].get_text()


@pytest.mark.skipif(os.environ.get('TEST_OFFICE') != '1', reason='Requires installed Microsoft Office')
@pytest.mark.parametrize('kind', ['Word', 'Excel', 'PowerPoint'])
def test_office_exports(kind, tmp_path):
    if kind == 'Word':
        from docx import Document
        source = tmp_path / 'sample.docx'
        doc = Document()
        doc.add_paragraph('Office conversion 12345')
        doc.save(source)
    elif kind == 'Excel':
        from openpyxl import Workbook
        source = tmp_path / 'sample.xlsx'
        doc = Workbook()
        doc.active['A1'] = 'Office conversion 12345'
        doc.save(source)
    else:
        from pptx import Presentation
        from pptx.util import Inches
        source = tmp_path / 'sample.pptx'
        doc = Presentation()
        slide = doc.slides.add_slide(doc.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = 'Office conversion 12345'
        doc.save(source)
    output = tmp_path / 'result.pdf'
    advanced.office_pdf(source, output, kind)
    with fitz.open(output) as pdf:
        assert '12345' in ''.join(p.get_text() for p in pdf)
