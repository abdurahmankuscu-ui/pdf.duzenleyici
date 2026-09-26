from pathlib import Path
import pymupdf as fitz
import pytest
from engine import Editor, add_text, erase, extract, page_range
import advanced


@pytest.fixture
def editor():
    e = Editor()
    e.new()
    yield e
    e.doc.close()


def test_turkish_save_redaction_undo(editor, tmp_path):
    rect = fitz.Rect(40, 40, 500, 110)
    editor.change(lambda d: add_text(d[0], rect, 'Türkçe: İstanbul, ışık, şeker, öğün', 16))
    assert 'İstanbul' in editor.doc[0].get_text()
    editor.change(lambda d: erase(d[0], rect))
    assert 'İstanbul' not in editor.doc[0].get_text()
    editor.undo()
    assert 'İstanbul' in editor.doc[0].get_text()
    editor.undo(True)
    path = tmp_path / 'sonuç.pdf'
    editor.save(path)
    with fitz.open(path) as doc:
        assert 'İstanbul' not in doc[0].get_text()
    assert not editor.dirty


def test_failed_edit_rolls_back(editor):
    before = editor.doc[0].get_text()
    with pytest.raises(ValueError):
        editor.change(lambda d: add_text(d[0], fitz.Rect(1, 1, 10, 10), 'very long content', 30))
    assert editor.doc[0].get_text() == before


def test_pages_merge_extract_reorder(editor, tmp_path):
    editor.change(lambda d: d.new_page())
    editor.change(lambda d: add_text(d[1], fitz.Rect(30, 30, 350, 100), 'SECOND'))
    editor.change(lambda d: d.select([1, 0]))
    assert 'SECOND' in editor.doc[0].get_text()
    path = tmp_path / 'split.pdf'
    extract(editor.doc, [0], path)
    with fitz.open(path) as part:
        assert len(part) == 1
        assert 'SECOND' in part[0].get_text()
        editor.change(lambda d: d.insert_pdf(part))
    assert len(editor.doc) == 3
    editor.change(lambda d: d.delete_page(1))
    assert len(editor.doc) == 2


def test_atomic_save_same_source_and_password(editor, tmp_path):
    path = tmp_path / 'test.pdf'
    editor.save(path, 'secret')
    with pytest.raises(ValueError):
        editor.open(path, 'wrong')
    editor.open(path, 'secret')
    editor.change(lambda d: d[0].set_rotation(90))
    editor.save(path)
    # Saving an opened protected PDF must not silently strip its password.
    with fitz.open(path) as doc:
        assert doc.needs_pass
        assert doc.authenticate('secret')
        assert doc[0].rotation == 90


def test_protected_pdf_content_survives_open_and_save(editor, tmp_path):
    path = tmp_path / 'test.pdf'
    editor.change(lambda d: d[0].insert_text((50, 60), 'Gizli içerik'))
    editor.save(path, 'secret')
    editor.open(path, 'secret')
    assert 'Gizli içerik' in editor.doc[0].get_text()
    editor.save(path)
    with fitz.open(path) as doc:
        doc.authenticate('secret')
        assert 'Gizli içerik' in doc[0].get_text()


def test_protected_pdf_keeps_password_on_save_as(editor, tmp_path):
    source = tmp_path / 'source.pdf'
    editor.save(source, 'secret')
    editor.open(source, 'secret')
    copy = tmp_path / 'copy.pdf'
    editor.save(copy)
    with fitz.open(copy) as doc:
        assert doc.needs_pass and doc.authenticate('secret')


def test_explicit_empty_password_removes_protection(editor, tmp_path):
    path = tmp_path / 'test.pdf'
    editor.save(path, 'secret')
    editor.open(path, 'secret')
    editor.save(path, '')
    with fitz.open(path) as doc:
        assert not doc.needs_pass
    editor.save(path)
    with fitz.open(path) as doc:
        assert not doc.needs_pass


def test_erase_removes_overlapping_annotations_links_and_fields(editor):
    page = editor.doc[0]
    page.insert_text((50, 60), 'TCKN 12345678901')
    secret = page.search_for('12345678901')[0]
    page.add_text_annot(secret.tl, 'TCKN 12345678901')
    page.add_text_annot((400, 400), 'unrelated note')
    page.insert_link({'kind': fitz.LINK_URI, 'from': secret, 'uri': 'https://example.com/12345678901'})
    widget = fitz.Widget()
    widget.field_name = 'tckn'
    widget.field_type = fitz.PDF_WIDGET_TYPE_TEXT
    widget.rect = secret
    widget.field_value = '12345678901'
    page.add_widget(widget)
    editor.change(lambda d: erase(d[0], secret))
    data = editor.doc.tobytes(garbage=4, deflate=True)
    with fitz.open(stream=data, filetype='pdf') as doc:
        page = doc[0]
        assert '12345678901' not in page.get_text()
        assert [a.info['content'] for a in page.annots(types=[fitz.PDF_ANNOT_TEXT])] == ['unrelated note']
        assert not page.get_links()
        assert not list(page.widgets())
        assert b'12345678901' not in fitz.open(stream=data, filetype='pdf').tobytes(expand=True)


def test_sanitize_removes_hidden_document_data(editor):
    from engine import sanitize
    editor.doc.set_metadata({'title': 'TCKN 12345678901', 'author': 'Deniz'})
    editor.doc.embfile_add('gizli.txt', b'12345678901')
    editor.doc[0].insert_text((50, 60), 'Visible')
    editor.change(sanitize)
    data = editor.doc.tobytes(garbage=4, deflate=True)
    with fitz.open(stream=data, filetype='pdf') as doc:
        assert not doc.metadata['title'] and not doc.metadata['author']
        assert doc.embfile_count() == 0
        assert 'Visible' in doc[0].get_text()
    editor.undo()
    assert editor.doc.metadata['title'] == 'TCKN 12345678901'


def test_ranges():
    assert page_range('1-3,5,2', 5) == [0, 1, 2, 4]
    for value in ['0', '6', '3-1', 'a', '', '1-2-3']:
        with pytest.raises(ValueError):
            page_range(value, 5)


def test_conversions_and_compare(editor, tmp_path):
    editor.change(lambda d: add_text(d[0], fitz.Rect(30, 30, 500, 100), 'Example report'))
    data = editor.doc.tobytes()
    advanced.pdf_word(data, tmp_path / 'report.docx')
    advanced.pdf_excel(data, tmp_path / 'report.xlsx')
    advanced.pdf_pptx(data, tmp_path / 'report.pptx')
    advanced.pdf_markdown(data, tmp_path / 'report.md')
    from docx import Document
    from openpyxl import load_workbook
    from pptx import Presentation
    assert 'Example report' in '\n'.join(p.text for p in Document(tmp_path / 'report.docx').paragraphs)
    assert 'Example report' in load_workbook(tmp_path / 'report.xlsx').active['A1'].value
    assert len(Presentation(tmp_path / 'report.pptx').slides) == 1
    assert 'Example report' in (tmp_path / 'report.md').read_text(encoding='utf-8')
    path = tmp_path / 'other.pdf'
    editor.save(path)
    assert 'fark bulunamadı' in advanced.compare(data, path)
    editor.change(lambda d: add_text(d[0], fitz.Rect(30, 150, 500, 200), 'Added'))
    assert 'Added' in advanced.compare(editor.doc.tobytes(), path)


def test_images_and_form(editor, tmp_path):
    from PIL import Image
    image = tmp_path / 'image.png'
    Image.new('RGB', (200, 100), 'red').save(image)
    output = tmp_path / 'images.pdf'
    advanced.images_pdf([image, image], output)
    with fitz.open(output) as doc:
        assert len(doc) == 2
    def insert(doc):
        page = doc[0]
        widget = fitz.Widget()
        widget.field_name = 'name'
        widget.field_type = fitz.PDF_WIDGET_TYPE_TEXT
        widget.rect = fitz.Rect(20, 20, 220, 60)
        widget.field_value = 'Test'
        page.add_widget(widget)
    editor.change(insert)
    editor.save(tmp_path / 'form.pdf')
    with fitz.open(tmp_path / 'form.pdf') as doc:
        page = doc[0]
        assert next(page.widgets()).field_value == 'Test'
